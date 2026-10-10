"""PDF → متن و Word روی **خروجیِ واقعیِ** Chrome و LibreOffice (`tests/fixtures/pdf`).

هر ادعا یکی از باگ‌هایی است که ابزارهای قبلی داشتند و روی همین فایل‌ها اندازه
گرفته شد: `pdftotext` «سلام» را «سالم» می‌داد، LibreOffice اصلاً PDF→DOCX نداشت، و
ستونِ فارسی از چپ خوانده می‌شد.
"""
from __future__ import annotations

import shutil
import subprocess

from app import pdftext as X
from app import pdftools as T
from tests.pdfgates import FIX, needs_ocr_fas, needs_pdf_text


async def _doc(tmp_path, name, *, images=False, ocr=0):
    src = await T.prepare(str(FIX / name), str(tmp_path))
    return src, await X.extract(src, str(tmp_path), max_pages=50, ocr_max_pages=ocr,
                                want_images=images)


@needs_pdf_text
async def test_poppler_alone_reverses_lam_alef(tmp_path):
    """کنترل: اگر روزی poppler این را درست کرد، توجیهِ کلِ `pdftext` باید بازبینی شود."""
    out = subprocess.run(["pdftotext", str(FIX / "fa_chrome.pdf"), "-"],
                         capture_output=True, text=True, timeout=60).stdout
    assert "سالم" in out and "سلام" not in out


@needs_pdf_text
async def test_chrome_persian_reads_correctly(tmp_path):
    _, doc = await _doc(tmp_path, "fa_chrome.pdf")
    txt = X.to_text(doc)
    for word in ("سلام دنیا", "بالا", "اسلام", "ملاقات", "پی‌دی‌اف", "(با پرانتز)", "۱۲۳۴",
                 "کلمه English در وسط"):
        assert word in txt, (word, txt)
    for broken in ("سالم", "باال", "مالقات", ")با پرانتز("):
        assert broken not in txt, (broken, txt)
    assert "An English paragraph with a فارسی word inside." in txt


@needs_pdf_text
async def test_libreoffice_zwnj_and_heading(tmp_path):
    _, doc = await _doc(tmp_path, "fa_lo.pdf")
    txt = X.to_text(doc)
    assert txt.startswith("گزارش ماهانه\n\n"), txt[:80]
    assert "پی‌دی‌اف به ورد" in txt
    assert "پی دی اف" not in txt


@needs_pdf_text
async def test_an_ltr_persian_paragraph_keeps_its_direction(tmp_path):
    """LibreOffice از TXT پاراگرافِ چپ‌به‌راست ساخته؛ متن درست و جهت نگه داشته شود."""
    _, doc = await _doc(tmp_path, "fa_lo_ltr.pdf")
    txt = X.to_text(doc)
    assert "این یک فایل متنی فارسی است." in txt
    para = next(p for p in doc.pages[0].paras if "فایل" in p.joined())
    assert para.rtl is False


@needs_pdf_text
async def test_persian_columns_are_read_right_first_english_left_first(tmp_path):
    _, doc = await _doc(tmp_path, "cols_chrome.pdf")
    paras = [p.joined() for p in doc.pages[0].paras]
    fa = [i for i, t in enumerate(paras) if X.has_rtl(t) and len(t) > 60]
    en = [i for i, t in enumerate(paras) if not X.has_rtl(t) and len(t) > 60]
    assert fa and en
    # راستِ نوارِ فارسی: پاراگرافی که بزرگ‌ترین x1 را دارد اول می‌آید
    fa_paras = [doc.pages[0].paras[i] for i in fa]
    assert fa_paras[0].x1 > fa_paras[-1].x1 + 50, [(p.x0, p.x1) for p in fa_paras]
    en_paras = [doc.pages[0].paras[i] for i in en]
    assert en_paras[0].x0 < en_paras[-1].x0 - 50, [(p.x0, p.x1) for p in en_paras]


@needs_pdf_text
async def test_word_output_is_editable_and_right_to_left(tmp_path):
    import docx
    from docx.oxml.ns import qn

    src, doc = await _doc(tmp_path, "fa_chrome.pdf", images=True)
    await X.render_images(src, str(tmp_path), doc)
    out = tmp_path / "o.docx"
    X.to_docx(doc, str(out))
    d = docx.Document(str(out))
    texts = [p.text for p in d.paragraphs if p.text]
    assert any("سلام دنیا" in t for t in texts)
    rtl = [p for p in d.paragraphs if "اسلام" in p.text]
    assert rtl and rtl[0]._p.pPr.find(qn("w:bidi")) is not None, "پاراگرافِ فارسی بی w:bidi"
    ltr = [p for p in d.paragraphs if p.text.startswith("An English")]
    assert ltr and (ltr[0]._p.pPr is None or ltr[0]._p.pPr.find(qn("w:bidi")) is None)
    # Word دیگر کادرِ شناورِ LibreOffice نیست: پاراگرافِ جاری
    assert not d.element.body.findall(".//" + qn("w:txbxContent"))


@needs_pdf_text
async def test_an_image_only_page_becomes_a_picture_in_word_without_ocr(tmp_path):
    png = await T.render_page(str(FIX / "fa_chrome.pdf"), 1, str(tmp_path / "pg"), dpi=100)
    scan = tmp_path / "scan.pdf"
    await T.images_to_pdf([png], str(scan), mode="fit", workdir=str(tmp_path))
    doc = await X.extract(str(scan), str(tmp_path), max_pages=5, ocr_max_pages=0, want_images=True)
    assert not X.has_text(doc) and doc.ocr_skipped == 1 and doc.pages[0].scan
    await X.render_images(str(scan), str(tmp_path), doc)
    assert len(doc.pages[0].images) == 1
    X.to_docx(doc, str(tmp_path / "o.docx"))
    import docx
    assert docx.Document(str(tmp_path / "o.docx")).inline_shapes, "صفحهٔ اسکن عکس نشد"


@needs_ocr_fas
async def test_a_scanned_page_is_read_with_ocr(tmp_path):
    png = await T.render_page(str(FIX / "fa_chrome.pdf"), 1, str(tmp_path / "pg"), dpi=200, gray=True)
    scan = tmp_path / "scan.pdf"
    await T.images_to_pdf([png], str(scan), mode="fit", workdir=str(tmp_path))
    doc = await X.extract(str(scan), str(tmp_path), max_pages=5, ocr_max_pages=5, want_images=False)
    assert doc.ocr_pages == 1
    txt = X.to_text(doc)
    # فقط کلمه‌هایی که tesseract پایدار می‌خواند — کیفیتِ خودِ OCR موضوعِ این تست نیست
    # («دنیا» را «Lod» خواند)، موضوع این است که صفحهٔ بی‌متن **اصلاً** خوانده شود.
    for word in ("گزارش", "آزمایشی", "English"):
        assert word in txt, (word, txt)


@needs_pdf_text
async def test_a_page_cap_is_reported_not_silent(tmp_path):
    many = tmp_path / "many.pdf"
    subprocess.run(["qpdf", "--empty", "--pages", *(str(FIX / "fa_chrome.pdf"),) * 3, "--", str(many)],
                   check=True, timeout=60)
    doc = await X.extract(str(many), str(tmp_path), max_pages=2, ocr_max_pages=0, want_images=False)
    assert doc.total == 3 and len(doc.pages) == 2


def test_text_file_to_word_sets_direction_per_line(tmp_path):
    import docx
    from docx.oxml.ns import qn

    out = tmp_path / "t.docx"
    X.text_to_docx("سلام دنیا\nاین یک فایل متنی فارسی است.\nSecond line English.\n", str(out))
    paras = [p for p in docx.Document(str(out)).paragraphs if p.text]
    bidi = [p._p.pPr is not None and p._p.pPr.find(qn("w:bidi")) is not None for p in paras]
    assert bidi == [True, True, False]


def test_an_old_windows_persian_text_file_is_decoded(tmp_path):
    p = tmp_path / "old.txt"
    p.write_bytes("متن قديمي با كاف".encode("cp1256"))
    assert X.read_text_file(str(p)) == "متن قديمي با كاف"
    u = tmp_path / "new.txt"
    u.write_bytes("﻿سلام".encode("utf-8"))
    assert X.read_text_file(str(u)) == "سلام"


def test_fixtures_are_real_tool_output():
    """فیکسچرِ دست‌ساز همان باگ‌ها را نمی‌سازد؛ تولیدکنندهٔ هر فایل ثابت بماند."""
    if not shutil.which("pdfinfo"):
        return
    for name, producer in (("fa_chrome.pdf", "Skia"), ("cols_chrome.pdf", "Skia"),
                           ("fa_lo.pdf", "LibreOffice"), ("fa_lo_ltr.pdf", "LibreOffice")):
        info = subprocess.run(["pdfinfo", str(FIX / name)], capture_output=True, text=True,
                              timeout=30).stdout
        assert producer in info, (name, info)
