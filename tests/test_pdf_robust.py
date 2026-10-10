"""PDF → Word/متن روی PDFهایی که «عادی» نیستند — همان‌هایی که تبدیل را می‌کشتند.

سه ردهٔ شکست، هر سه روی کدِ پیش از رفع بازتولید شد:

* **نویسهٔ کنترلی از قلم.** قلمی که ToUnicodeش گلیف را به کدِ کنترلی نگاشته (رایج‌ترینش
  ‎\\x03‎: شمارهٔ گلیفِ فاصله در قلم‌های TrueType). poppler آن بایت را **خام** در XMLِ
  `-bbox-layout` می‌نویسد و expat با `not well-formed` کلِ تبدیل را می‌کشت — روی صفحهٔ
  انگلیسی هم؛ همان نویسه یا NUL/U+FFFE از pdfplumber به python-docx می‌رسید و lxml با
  `All strings must be XML compatible` رد می‌کرد.
* **صفحه‌ای که اندازه‌اش پیکسل است.** اسکنِ گوشی/PILِ پیش‌فرض صفحهٔ ۲۵۵۰×۳۳۰۰ pt می‌سازد؛
  OCR در ۳۰۰dpi یعنی تصویرِ ۱۴۶ مگاپیکسلی، ~۱٫۶ گیگابایت حافظهٔ tesseract و ده‌ها ثانیه
  برای **یک** صفحه — و دقتِ کمتر (عکسِ مبدأ خودش ۲۵۵۰ پیکسل است).
* **صفحهٔ Wordِ بیرون از حدِ Word** (۲۲ اینچ) برای همان اسکن‌ها.

PDFهای «قلمِ خراب» این‌جا در خودِ تست ساخته می‌شوند (بایت‌به‌بایت، بی‌کتابخانه)، چون
نکته دقیقاً یک جدولِ ToUnicodeِ مشخص است؛ کنترلِ `test_the_harness_reproduces_the_raw_byte`
ثابت می‌کند poppler روی همین فایل همان XMLِ خراب را می‌سازد.
"""
from __future__ import annotations

import os
import re
import subprocess
import types
import xml.etree.ElementTree as ET
import zipfile

import pytest

from app import pdftext as X
from app import pdftools as T
from app import tasks
from tests.pdfgates import FIX, needs_pdf_text, needs_pdf_tools

#: الف، لام، میم، فاصله — «ABC» روی صفحه یعنی ترتیبِ دیداریِ «ملا»
_FA = [(0x41, "0627"), (0x42, "0644"), (0x43, "0645"), (0x20, "0020")]
_LAT = [(0x41, "0048"), (0x42, "0069"), (0x43, "0021"), (0x20, "0020")]


def _bad_font_pdf(path, cmap: list[tuple[int, str]], text: bytes = b"ABC ABDC") -> str:
    """یک صفحه با قلمی که ToUnicodeش `cmap` است (`D` = نویسهٔ مشکل‌دار). نامِ قلم عمداً
    از ۱۴ قلمِ استاندارد نیست، وگرنه pdfminer `/Widths` را نادیده می‌گیرد و همهٔ گلیف‌ها
    پهنای صفر می‌گیرند (اندازه‌گیری‌شده)."""
    bf = "\n".join(f"<{s:02X}> <{d}>" for s, d in cmap)
    tounicode = (b"/CIDInit /ProcSet findresource begin 12 dict begin begincmap\n"
                 b"/CMapName /X def /CMapType 2 def\n1 begincodespacerange <00> <FF> "
                 b"endcodespacerange\n" + f"{len(cmap)} beginbfchar\n{bf}\nendbfchar\n".encode()
                 + b"endcmap CMapName currentdict /CMap defineresource pop end end")
    content = b"BT /F1 24 Tf 72 700 Td (" + text + b") Tj ET"
    widths = b" ".join([b"300"] + [b"600"] * 36)
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
            b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
            b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /TlzTest /FirstChar 32 /LastChar 68 "
            b"/Widths [" + widths + b"] /ToUnicode 6 0 R >>",
            b"<< /Length %d >>\nstream\n" % len(tounicode) + tounicode + b"\nendstream"]
    out, offs = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offs.append(len(out))
        out += b"%d 0 obj\n" % i + o + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offs)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)
    p = os.path.join(str(path), "bad-font.pdf")
    with open(p, "wb") as fh:
        fh.write(out)
    return p


def _phone_scan(tmp_path, dpi: int = 300) -> str:
    """اسکنِ «گوشی»: صفحهٔ واقعیِ فارسی در ۳۰۰dpi به عکس، و عکس با PILِ پیش‌فرض (۷۲dpi) به
    PDF — صفحه ۲۵۵۰×۳۳۰۰ pt می‌شود، یعنی اندازه‌اش پیکسل است نه کاغذ."""
    from PIL import Image

    base = str(tmp_path / "shot")
    subprocess.run(["pdftoppm", "-r", str(dpi), "-png", "-singlefile",
                    str(FIX / "fa_chrome.pdf"), base], check=True, timeout=120)
    out = str(tmp_path / "phone.pdf")
    with Image.open(base + ".png") as im:
        im.convert("RGB").save(out)
    return out


class PathBot:
    async def get_file(self, fid):
        return types.SimpleNamespace(file_path=fid)


async def _to(tmp_path, src, fmt):
    w = tmp_path / f"w-{fmt}-{len(os.listdir(tmp_path))}"
    w.mkdir()
    f = types.SimpleNamespace(name="doc.pdf", kind="pdf", mime="application/pdf",
                              duration=None, width=None, height=None, size=None)
    return await tasks._do_op(PathBot(), "convert", {"target": fmt}, f, src, str(w), "fa",
                              redis=None)


def _docx_xml(path) -> str:
    with zipfile.ZipFile(path) as z:
        return z.read("word/document.xml").decode("utf-8")


def _docx_text(path) -> str:
    import docx

    return "\n".join(p.text for p in docx.Document(path).paragraphs)


# ── نویسهٔ کنترلی از قلم ─────────────────────────────────────────
@needs_pdf_text
def test_the_harness_reproduces_the_raw_byte(tmp_path):
    """کنترل: poppler روی همین فایل بایتِ ‎\\x03‎ را خام می‌نویسد و expat ردش می‌کند —
    یعنی تست‌های زیر همان شکستِ تولید را می‌سنجند، نه یک فایلِ بی‌خطر."""
    src = _bad_font_pdf(tmp_path, _LAT + [(0x44, "0003")])
    out = subprocess.run(["pdftotext", "-bbox-layout", "-enc", "UTF-8", src, "-"],
                         capture_output=True, timeout=60).stdout
    assert b"Hi\x03!" in out
    with pytest.raises(ET.ParseError):
        ET.fromstring(out)


@needs_pdf_text
async def test_a_control_code_from_the_font_does_not_kill_an_english_page(tmp_path):
    src = _bad_font_pdf(tmp_path, _LAT + [(0x44, "0003")])
    doc = await X.extract(src, str(tmp_path), max_pages=5, ocr_max_pages=0, want_images=False)
    # ‎\x03‎ پشتِ یک گلیفِ جادار است، پس فاصله می‌شود نه چسباندنِ دو تکه
    assert X.to_text(doc) == "Hi! Hi !\n"


@needs_pdf_text
@pytest.mark.parametrize("bad", ["0003", "0000", "FFFE"], ids=["ctrl3", "nul", "fffe"])
async def test_word_from_a_persian_page_with_a_bad_glyph(tmp_path, bad):
    """صفحهٔ فارسی از مسیرِ گلیف‌های pdfplumber می‌رود — همان مسیری که NUL و U+FFFE را
    به python-docx می‌رساند."""
    src = _bad_font_pdf(tmp_path, _FA + [(0x44, bad)])
    res = await _to(tmp_path, src, "docx")
    path = res["spawn"]["path"]
    text = _docx_text(path)
    assert "ملا" in text                         # کلمهٔ سالم سالم ماند
    assert not re.search("[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]", _docx_xml(path))


@needs_pdf_text
async def test_text_output_carries_no_control_characters(tmp_path):
    src = _bad_font_pdf(tmp_path, _FA + [(0x44, "0000")])
    res = await _to(tmp_path, src, "txt")
    with open(res["spawn"]["path"], encoding="utf-8") as fh:
        body = fh.read()
    assert "ملا" in body and "\x00" not in body


@needs_pdf_text
def test_the_word_writer_is_the_last_guard(tmp_path):
    """لایهٔ دوم، جدا از اولی (§۶ «دفاعِ لایه‌ای آزمونِ لایه‌ای می‌خواهد»): متنی که از
    هیچ پاک‌کننده‌ای رد نشده — این‌جا OCR — باز هم سندِ Word را نمی‌کشد."""
    page = X.Page(no=1, width=595, height=842, ocr="متن\x00 OCR\ufffe پایان")
    out = str(tmp_path / "g.docx")
    X.to_docx(X.Doc(pages=[page], total=1), out)
    assert "متن OCR پایان" in _docx_text(out)


@needs_pdf_text
def test_xml_safe_removes_exactly_what_word_rejects():
    """کشف‌محور: هر نویسه‌ای که خودِ python-docx/lxml رد می‌کند باید برداشته شود، و هر
    نویسه‌ای که می‌پذیرد (تب، خطِ جدید، حروف، نیم‌فاصله) بماند — نه یک فهرستِ دست‌نویسِ
    دوم که با lxml واگرا شود."""
    import docx

    candidates = [*range(0x20), 0x7F, 0x85, 0x200C, 0x0627, 0xFFFD, 0xFFFE, 0xFFFF,
                  0xD800, 0xDBFF, 0xDC00, 0xDFFF]
    for cp in candidates:
        ch = chr(cp)
        try:
            docx.Document().add_paragraph().add_run(ch)
            rejected = False
        except (ValueError, UnicodeEncodeError):
            rejected = True
        assert (X.xml_safe(ch) == "") == rejected, hex(cp)
    assert sum(X.xml_safe(chr(cp)) == "" for cp in candidates) >= 30   # ضدِتوخالی


@needs_pdf_text
def test_a_text_file_with_a_form_feed_and_nul_becomes_word(tmp_path):
    """مسیرِ «متن → PDF» هم از python-docx رد می‌شود؛ فرم‌فیدِ فایلِ متنیِ قدیمی شکستِ
    سطر است."""
    out = str(tmp_path / "t.docx")
    X.text_to_docx("سطر اول\x0cسطر دوم\x00 تمام", out)
    paras = [p for p in _docx_text(out).split("\n") if p]
    assert paras == ["سطر اول", "سطر دوم تمام"]


# ── صفحه‌ای که اندازه‌اش پیکسل است ───────────────────────────────
def test_fit_dpi():
    assert T.fit_dpi((595, 842), 300, 4200) == 300          # A4 همان ۳۰۰dpi
    assert T.fit_dpi((2550, 3300), 300, 4200) == 91         # اسکنِ پیکسلی
    assert T.fit_dpi((3300, 2550), 300, 4200) == 91         # افقی هم
    assert T.fit_dpi(None, 300, 4200) == 300                # اندازه نامعلوم
    assert T.fit_dpi((0, 0), 150, 10) == 150
    assert T.fit_dpi((10**6, 10**6), 300, 4200) == 1        # هرگز صفر


async def _ocr_sizes(monkeypatch, src, workdir):
    """اندازهٔ تصویری که به tesseract می‌رسد — بی‌خودِ tesseract."""
    from PIL import Image

    seen: list[tuple[int, int]] = []

    async def fake_ocr(png, cancel=None, langs="fas+eng"):
        with Image.open(png) as im:
            seen.append(im.size)
        return "متن"

    monkeypatch.setattr(T, "ocr_png", fake_ocr)
    monkeypatch.setattr(T, "have", lambda tool: True)
    await T.ocr_pages(src, workdir, [1])
    return seen


@needs_pdf_tools
async def test_a_pixel_sized_page_is_ocred_at_a_capped_size(tmp_path, monkeypatch):
    src = _phone_scan(tmp_path)
    (w,) = await _ocr_sizes(monkeypatch, src, str(tmp_path))
    assert max(w) <= T.OCR_MAX_PX            # پیش از رفع ۱۰۶۲۵×۱۳۷۵۰ (۱۴۶ مگاپیکسل)


@needs_pdf_tools
async def test_an_ordinary_page_keeps_full_ocr_resolution(tmp_path, monkeypatch):
    """کنترل: سقف فقط صفحهٔ غول‌پیکر را می‌گیرد؛ Letter همچنان ۳۰۰dpi است."""
    (w,) = await _ocr_sizes(monkeypatch, str(FIX / "fa_chrome.pdf"), str(tmp_path))
    assert max(w) == 3300                    # ۷۹۲ pt × ۳۰۰/۷۲


@needs_pdf_tools
async def test_word_pictures_are_capped_too(tmp_path):
    src = _phone_scan(tmp_path, dpi=150)
    doc = await X.extract(src, str(tmp_path), max_pages=5, ocr_max_pages=0, want_images=True)
    assert doc.pages[0].scan                 # بی‌OCR → عکسِ تمام‌صفحه
    await X.render_images(src, str(tmp_path), doc)
    from PIL import Image

    with Image.open(doc.pages[0].images[0].path) as im:
        assert max(im.size) <= X.IMG_MAX_PX


# ── یک صفحهٔ OCRِ شکسته ──────────────────────────────────────────
@needs_pdf_tools
async def test_one_failed_ocr_page_does_not_fail_the_rest(tmp_path, monkeypatch):
    calls = {"n": 0}

    async def flaky(png, cancel=None, langs="fas+eng"):
        calls["n"] += 1
        if calls["n"] == 1:
            raise RuntimeError("tesseract failed (code -9)")      # OOM-kill
        return "صفحهٔ دوم"

    monkeypatch.setattr(T, "ocr_png", flaky)
    monkeypatch.setattr(T, "have", lambda tool: True)
    two = str(tmp_path / "two.pdf")
    subprocess.run(["qpdf", "--empty", "--pages", str(FIX / "fa_chrome.pdf"),
                    str(FIX / "fa_lo.pdf"), "--", two], check=True, timeout=60)
    assert await T.ocr_pages(two, str(tmp_path), [1, 2]) == [None, "صفحهٔ دوم"]


@needs_pdf_tools
async def test_ocr_that_fails_everywhere_is_still_an_error(tmp_path, monkeypatch):
    """کنترل: tesseractِ خراب «این PDF متن ندارد» نشان داده نشود."""
    async def broken(png, cancel=None, langs="fas+eng"):
        raise RuntimeError("tesseract failed (code 1)")

    monkeypatch.setattr(T, "ocr_png", broken)
    monkeypatch.setattr(T, "have", lambda tool: True)
    with pytest.raises(RuntimeError, match="tesseract failed"):
        await T.ocr_pages(str(FIX / "fa_chrome.pdf"), str(tmp_path), [1])


@needs_pdf_text
async def test_a_page_whose_ocr_failed_becomes_a_picture_in_word(tmp_path, monkeypatch):
    async def fail(*a, **kw):
        return [None]

    monkeypatch.setattr(T, "ocr_pages", fail)
    monkeypatch.setattr(T, "have", lambda tool: True)
    src = _phone_scan(tmp_path, dpi=100)
    doc = await X.extract(src, str(tmp_path), max_pages=5, ocr_max_pages=5, want_images=True)
    assert doc.ocr_pages == 0 and doc.pages[0].scan


# ── اندازهٔ صفحهٔ Word ───────────────────────────────────────────
@needs_pdf_text
async def test_a_pixel_sized_page_makes_a_word_page_word_can_open(tmp_path):
    src = _phone_scan(tmp_path, dpi=150)
    doc = await X.extract(src, str(tmp_path), max_pages=5, ocr_max_pages=0, want_images=True)
    await X.render_images(src, str(tmp_path), doc)
    out = str(tmp_path / "p.docx")
    X.to_docx(doc, out)
    sz = re.search(r'<w:pgSz w:w="(\d+)" w:h="(\d+)"', _docx_xml(out))
    w, h = int(sz[1]), int(sz[2])
    assert max(w, h) <= 22 * 1440            # حدِ Word؛ پیش از رفع ۲۵۵۰۰×۳۳۰۰۰
    assert abs(h - 16838) <= 2               # ضلعِ بلند = A4
    assert abs(w / h - 1275 / 1650) < 0.01   # نسبتِ صفحه همان


@needs_pdf_text
async def test_an_ordinary_page_keeps_its_size_in_word(tmp_path):
    """کنترل: Letterِ Chrome همان ۶۱۲×۷۹۲ pt می‌ماند."""
    src = str(FIX / "fa_chrome.pdf")
    doc = await X.extract(src, str(tmp_path), max_pages=5, ocr_max_pages=0, want_images=False)
    out = str(tmp_path / "l.docx")
    X.to_docx(doc, out)
    assert '<w:pgSz w:w="12240" w:h="15840"' in _docx_xml(out)


def test_word_page_and_font_limits():
    assert X.page_scale(612, 792) == 1.0
    assert X.page_scale(22 * 72, 22 * 72) == 1.0             # درست روی مرز
    assert X.page_scale(2550, 3300) == pytest.approx(841.89 / 3300)
    assert X._font_pt(0.1) == 1.0 and X._font_pt(5000) == 1638.0 and X._font_pt(11.3) == 11.5
