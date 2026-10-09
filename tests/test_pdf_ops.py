"""عملیاتِ PDF از درِ ورکر (`tasks._do_op`/`run_op`) و درِ ربات (`routers/ops`).

`_do_op` با ابزارهای واقعی و PDFهای واقعیِ `tests/fixtures/pdf` اجرا می‌شود؛ تنها
چیزِ جعلی بات است (همان `ValidatingBot` که فراخوانی را با امضای aiogram می‌سنجد).
"""
from __future__ import annotations

import asyncio
import os
import subprocess
import sys
import types
from datetime import datetime, timezone

import pytest
import pytest_asyncio
from aiogram.fsm.context import FSMContext
from aiogram.fsm.storage.base import StorageKey
from aiogram.fsm.storage.memory import MemoryStorage
from aiogram.types import Chat, Document, Message
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from app import pagespec
from app import processing as P
from app import tasks as T
from app.exceptions import ProcessingCancelled, UserFacingError
from app.i18n import t
from app.models import Base, File, Job, User
from app.routers import ops
from app.states import Collect, PdfPages, PdfPassword
from tests.aiogram_double import ValidatingBot, bind_like_aiogram
from tests.pdfgates import FIX, needs_pdf_text, needs_pdf_tools
from tests.test_probe_orphan import _alive

FA = str(FIX / "fa_chrome.pdf")
CHAT = 4242


class PathBot:
    """`get_file` مسیرِ محلی را برمی‌گرداند — همان شکلِ مستر (`is_local=True`)."""

    async def get_file(self, fid):
        return types.SimpleNamespace(file_path=fid)


def _file(name="doc.pdf", kind="pdf", mime="application/pdf"):
    return types.SimpleNamespace(name=name, kind=kind, mime=mime, duration=None,
                                 width=None, height=None, size=None)


async def _op(tmp_path, op, args, src=FA, *, name="doc.pdf", kind="pdf",
              mime="application/pdf", redis=None):
    w = tmp_path / f"w-{op}-{len(os.listdir(tmp_path))}"
    w.mkdir()
    return await T._do_op(PathBot(), op, args, _file(name, kind, mime), src, str(w), "fa",
                          redis=redis)


def _three(tmp_path) -> str:
    out = tmp_path / "three.pdf"
    subprocess.run(["qpdf", "--empty", "--pages", FA, FIX / "fa_lo.pdf",
                    FIX / "cols_chrome.pdf", "--", out], check=True, timeout=60)
    return str(out)


def _locked(tmp_path, pw="سری") -> str:
    out = tmp_path / "locked.pdf"
    subprocess.run(["qpdf", "--encrypt", pw, "owner-x", "256", "--", FA, out],
                   check=True, timeout=60)
    return str(out)


# ── تبدیل ──────────────────────────────────────────────────────
@needs_pdf_text
async def test_word_and_text_are_new_cards_and_the_pdf_stays(tmp_path):
    res = await _op(tmp_path, "convert", {"target": "docx"}, name="گزارش.pdf")
    assert "path" not in res and res["spawn"]["kind"] == "document"
    assert res["spawn"]["name"] == "گزارش.docx" and os.path.getsize(res["spawn"]["path"]) > 1000
    res = await _op(tmp_path, "convert", {"target": "txt"})
    text = open(res["spawn"]["path"], encoding="utf-8").read()
    assert "سلام دنیا" in text and "سالم" not in text


@needs_pdf_text
async def test_page_images_come_as_an_album_or_a_zip(tmp_path, monkeypatch):
    src = _three(tmp_path)
    res = await _op(tmp_path, "convert", {"target": "jpg"}, src, name="book.pdf")
    assert res["album"] is True
    assert [os.path.basename(p) for p in res["files"]] == ["book-1.jpg", "book-2.jpg", "book-3.jpg"]
    monkeypatch.setattr(T.pdftools, "ALBUM_MAX_FILES", 2)
    res = await _op(tmp_path, "convert", {"target": "png"}, src, name="book.pdf")
    assert len(res["files"]) == 1 and res["files"][0].endswith("book-pages.zip")
    assert not res.get("album")


@needs_pdf_text
async def test_a_locked_pdf_says_so_instead_of_a_tool_error(tmp_path):
    with pytest.raises(UserFacingError) as e:
        await _op(tmp_path, "convert", {"target": "txt"}, _locked(tmp_path))
    assert e.value.key == "pdf_needs_password"


# ── کاهشِ حجم ──────────────────────────────────────────────────
@needs_pdf_tools
async def test_compress_replaces_in_place_or_says_no_gain(tmp_path, monkeypatch):
    from tests.test_pdftools import _scan_jpeg
    scan = tmp_path / "scan.pdf"
    await T.pdftools.images_to_pdf([_scan_jpeg(tmp_path / "a.jpg")], str(scan), workdir=str(tmp_path))
    res = await _op(tmp_path, "compress", {"level": "strong"}, str(scan), name="scan.pdf")
    assert res["kind"] == "pdf" and res["filename"] == "scan-min.pdf"
    assert os.path.getsize(res["path"]) < os.path.getsize(scan)
    monkeypatch.setattr(T.pdftools, "MIN_GAIN", 0.0)
    res = await _op(tmp_path, "compress", {"level": "normal"}, str(scan))
    assert res == {"note_only": True, "label": t("fa", "cl_pdf_no_gain")}


# ── صفحه‌ها ───────────────────────────────────────────────────
@needs_pdf_tools
async def test_extract_is_a_new_card_delete_edits_in_place(tmp_path):
    src = _three(tmp_path)
    res = await _op(tmp_path, "pdf_select", {"spec": "3, 1", "mode": "keep"}, src, name="b.pdf")
    assert res["spawn"]["kind"] == "pdf" and res["spawn"]["name"] == "b-p3_1.pdf"
    res = await _op(tmp_path, "pdf_select", {"spec": "۲", "mode": "delete"}, src, name="b.pdf")
    assert res["filename"] == "b.pdf" and "spawn" not in res
    n = subprocess.run(["qpdf", "--show-npages", res["path"]], capture_output=True, text=True).stdout
    assert int(n) == 2


@needs_pdf_tools
@pytest.mark.parametrize("args,key", [
    ({"spec": "9", "mode": "keep"}, "pdf_page_out_of_range"),
    ({"spec": "1-3", "mode": "delete"}, "pdf_delete_all"),
    ({"spec": "x", "mode": "keep"}, "pdf_pages_bad"),
], ids=["out-of-range", "delete-all", "garbage"])
async def test_page_errors_are_user_messages(tmp_path, args, key):
    with pytest.raises(UserFacingError) as e:
        await _op(tmp_path, "pdf_select", args, _three(tmp_path))
    assert e.value.key == key
    if key == "pdf_page_out_of_range":
        assert e.value.kw == {"page": 9, "n": 3}


@needs_pdf_tools
async def test_rotate_and_split(tmp_path):
    res = await _op(tmp_path, "pdf_rotate", {"angle": 270}, FA)
    assert res["kind"] == "pdf"
    with pytest.raises(ValueError):
        await _op(tmp_path, "pdf_rotate", {"angle": 45}, FA)
    with pytest.raises(UserFacingError) as e:
        await _op(tmp_path, "pdf_split", {}, FA)
    assert e.value.key == "pdf_split_single"
    res = await _op(tmp_path, "pdf_split", {}, _three(tmp_path), name="b.pdf")
    assert res["album"] and len(res["files"]) == 3


# ── رمز ────────────────────────────────────────────────────────
@needs_pdf_tools
async def test_lock_then_unlock_and_the_password_is_used_once(tmp_path, redis):
    key = pagespec.PW_KEY.format(tok="t1")
    await redis.set(key, "رمز۱", ex=60)
    res = await _op(tmp_path, "pdf_lock", {"tok": "t1"}, FA, redis=redis)
    assert await redis.get(key) is None, "رمز بعد از جاب در Redis ماند"
    locked = res["path"]
    with pytest.raises(UserFacingError) as e:      # همان توکن دوباره → منقضی
        await _op(tmp_path, "pdf_lock", {"tok": "t1"}, FA, redis=redis)
    assert e.value.key == "pdf_pw_expired"

    await redis.set(pagespec.PW_KEY.format(tok="t2"), "غلط", ex=60)
    with pytest.raises(UserFacingError) as e:
        await _op(tmp_path, "pdf_unlock", {"tok": "t2"}, locked, redis=redis)
    assert e.value.key == "pdf_wrong_password"

    await redis.set(pagespec.PW_KEY.format(tok="t3"), "رمز۱", ex=60)
    res = await _op(tmp_path, "pdf_unlock", {"tok": "t3"}, locked, redis=redis)
    assert res["kind"] == "pdf"
    assert not await T.pdftools.is_encrypted(res["path"])


@needs_pdf_tools
async def test_unlock_without_a_user_password(tmp_path, redis):
    """فقط قفلِ مالک: هر رمزی کافی است و قفل برداشته می‌شود؛ بی‌قفل: «رمز نداشت»."""
    owner = tmp_path / "owner.pdf"
    subprocess.run(["qpdf", "--encrypt", "", "own", "256", "--", FA, owner], check=True)
    await redis.set(pagespec.PW_KEY.format(tok="a"), "هرچی", ex=60)
    res = await _op(tmp_path, "pdf_unlock", {"tok": "a"}, str(owner), redis=redis)
    assert res["label"] == t("fa", "cl_pdf_unlocked")
    await redis.set(pagespec.PW_KEY.format(tok="b"), "هرچی", ex=60)
    res = await _op(tmp_path, "pdf_unlock", {"tok": "b"}, FA, redis=redis)
    assert res == {"note_only": True, "label": t("fa", "cl_pdf_not_locked")}


# ── ادغام و عکس‌ها ─────────────────────────────────────────────
@needs_pdf_tools
async def test_merge_is_a_new_card_and_names_a_locked_member(tmp_path):
    members = [{"file_id": FA, "name": "a.pdf"}, {"file_id": str(FIX / "fa_lo.pdf"), "name": "b.pdf"}]
    res = await _op(tmp_path, "pdf_merge", {"members": members}, FA, name="a.pdf")
    assert res["spawn"]["name"] == "a-merged.pdf" and res["spawn"]["kind"] == "pdf"
    bad = [members[0], {"file_id": _locked(tmp_path), "name": "<b>x</b>.pdf"}]
    with pytest.raises(UserFacingError) as e:
        await _op(tmp_path, "pdf_merge", {"members": bad}, FA)
    assert e.value.key == "pdf_needs_password_member"
    assert e.value.kw["name"] == "&lt;b&gt;x&lt;/b&gt;.pdf", "نامِ کاربر در پیامِ HTML escape نشد"


@needs_pdf_tools
async def test_images_to_pdf_modes(tmp_path):
    from PIL import Image
    im = tmp_path / "p.png"
    Image.new("RGB", (300, 150), "red").save(im)
    m = [{"file_id": str(im)}]
    for mode, size in (("a4", "841.89 x 595.276"), ("fit", "144 x 72")):
        res = await _op(tmp_path, "images_to_pdf", {"members": m, "mode": mode}, str(im),
                        name="p.png", kind="image", mime="image/png")
        assert res["spawn"]["kind"] == "pdf"
        info = subprocess.run(["pdfinfo", res["spawn"]["path"]], capture_output=True, text=True).stdout
        assert size in info, (mode, info)


# ── سند → PDF (بدونِ LibreOfficeِ واقعی: یک `soffice`ِ ساختگی روی PATH) ──
_FAKE_SOFFICE = r"""
import os, shutil, sys
args = sys.argv[1:]
src = args[-1]
outdir = args[args.index("--outdir") + 1]
open(os.environ["FAKE_STATE"] + ".in", "w").write(src)
shutil.copyfile(os.environ["FAKE_PDF"], os.path.join(outdir, os.path.splitext(os.path.basename(src))[0] + ".pdf"))
"""


@pytest.fixture
def fake_soffice(tmp_path, monkeypatch):
    bindir = tmp_path / "bin"
    bindir.mkdir()
    p = bindir / "soffice"
    p.write_text(f"#!{sys.executable}\n" + _FAKE_SOFFICE)
    p.chmod(0o755)
    monkeypatch.setenv("PATH", f"{bindir}:{os.environ['PATH']}")
    monkeypatch.setenv("FAKE_STATE", str(tmp_path / "so"))
    monkeypatch.setenv("FAKE_PDF", FA)
    return tmp_path / "so.in"


async def test_document_to_pdf_is_a_pdf_card(tmp_path, fake_soffice):
    doc = tmp_path / "in.bin"
    doc.write_bytes(b"PK\x03\x04 not really a docx")
    res = await _op(tmp_path, "to_pdf", {}, str(doc), name="report",
                    kind="document", mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert res["spawn"]["kind"] == "pdf", "پیش از این `document` بود و منوی PDF نمی‌آمد"
    assert fake_soffice.read_text().endswith("report.docx"), "پسوند از mime نیامد"


async def test_a_text_file_goes_through_word_for_direction(tmp_path, fake_soffice):
    pytest.importorskip("docx")
    txt = tmp_path / "n.txt"
    txt.write_text("این یک فایل متنی فارسی است.\n", encoding="utf-8")
    await _op(tmp_path, "to_pdf", {}, str(txt), name="n.txt", kind="document", mime="text/plain")
    fed = fake_soffice.read_text()
    assert fed.endswith("n.docx"), fed
    import docx
    from docx.oxml.ns import qn
    p = [x for x in docx.Document(fed).paragraphs if x.text][0]
    assert p._p.pPr.find(qn("w:bidi")) is not None


@pytest.mark.parametrize("name,mime,key", [
    ("tool.exe", "application/x-msdownload", "to_pdf_unsupported"),
    ("a.<b>", "application/octet-stream", "to_pdf_unsupported"),
    ("empty.txt", "text/plain", "to_pdf_empty"),
], ids=["exe", "html-in-ext", "empty-text"])
async def test_document_to_pdf_refusals(tmp_path, name, mime, key):
    f = tmp_path / "f"
    f.write_text("  \n")
    with pytest.raises(UserFacingError) as e:
        await _op(tmp_path, "to_pdf", {}, str(f), name=name, kind="document", mime=mime)
    assert e.value.key == key
    if name == "a.<b>":
        assert "<" not in e.value.kw["ext"]


# ── LibreOffice یک درخت است: لغو باید کلِ گروه را بکشد ──────────────
_TREE = r"""
import os, subprocess, sys, time
here = os.environ["FAKE_STATE"]
child = subprocess.Popen([sys.executable, "-c",
    "import os,time\nopen(os.environ['FAKE_STATE']+'.child','w').write(str(os.getpid()))\n"
    "while True: time.sleep(0.05)"])
open(here + ".pid", "w").write(str(os.getpid()))
while True:
    time.sleep(0.05)
"""


async def _read_pid(path) -> int:
    for _ in range(200):
        try:
            return int(open(path).read())
        except (FileNotFoundError, ValueError):
            await asyncio.sleep(0.02)
    raise AssertionError(f"{path} نوشته نشد — فرایندِ جعلی اجرا نشد")


async def _gone(pid: int) -> bool:
    for _ in range(100):
        if not _alive(pid):
            return True
        await asyncio.sleep(0.05)
    return False


@pytest.mark.parametrize("how", ["job-cancel", "cancel-button"])
async def test_office_convert_kills_the_whole_tree(tmp_path, monkeypatch, how):
    """`soffice` → `soffice.bin`: کشتنِ اسکریپت فرزند را یتیم می‌گذاشت (اندازه‌گیری‌شده)."""
    bindir = tmp_path / "bin"
    bindir.mkdir()
    p = bindir / "soffice"
    p.write_text(f"#!{sys.executable}\n" + _TREE)
    p.chmod(0o755)
    monkeypatch.setenv("PATH", f"{bindir}:{os.environ['PATH']}")
    state = tmp_path / "tree"
    monkeypatch.setenv("FAKE_STATE", str(state))
    monkeypatch.setattr(P, "_CANCEL_POLL", 0.05)
    flag = {"v": False}

    async def cancel():
        return flag["v"]

    task = asyncio.create_task(P.office_convert(str(tmp_path / "x.docx"), str(tmp_path), "pdf",
                                                cancel=cancel))
    top = await _read_pid(str(state) + ".pid")
    child = await _read_pid(str(state) + ".child")
    try:
        if how == "job-cancel":
            task.cancel()
            with pytest.raises(asyncio.CancelledError):
                await asyncio.wait_for(task, 5)
        else:
            flag["v"] = True
            with pytest.raises(ProcessingCancelled):
                await asyncio.wait_for(task, 5)
        assert await _gone(top), "اسکریپتِ soffice زنده ماند"
        assert await _gone(child), "soffice.bin (فرزند) یتیم ماند"
    finally:
        for pid in (top, child):
            if _alive(pid):
                os.kill(pid, 9)


# ── run_op: پیامِ کاربر و کلیدِ پایدار ─────────────────────────────
class CapBot(ValidatingBot):
    def __init__(self) -> None:
        self.captions: list[str] = []
        self.calls: list[tuple[str, dict]] = []

    def _on(self, name, payload):
        self.calls.append((name, payload))
        if name == "edit_message_caption" and payload.get("caption"):
            self.captions.append(payload["caption"])
        return True


@pytest_asyncio.fixture
async def job_env(monkeypatch, tmp_path):
    monkeypatch.setattr(T.settings, "work_dir", str(tmp_path / "work"))
    engine = create_async_engine("sqlite+aiosqlite:///:memory:")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    maker = async_sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)
    monkeypatch.setattr(T, "Sessionmaker", maker)

    async def _noop(*a, **kw):
        return None
    monkeypatch.setattr(T.textstore, "refresh_if_stale", _noop)
    yield maker
    await engine.dispose()


async def _seed(maker, src, op, args) -> int:
    async with maker() as s:
        u = User(tg_user_id=5, role="user")
        s.add(u)
        await s.flush()
        f = File(ref="PdfJob01", owner_id=u.id, file_unique_id="u", file_id=src,
                 name="locked.pdf", kind="pdf", size=os.path.getsize(src), changelog=[])
        s.add(f)
        await s.flush()
        j = Job(file_id=f.id, op=op, args=args, status="queued")
        s.add(j)
        await s.commit()
        return j.id


@needs_pdf_tools
async def test_run_op_shows_the_users_message_and_stores_a_stable_key(tmp_path, job_env):
    maker = job_env
    job_id = await _seed(maker, _locked(tmp_path), "convert", {"target": "txt"})
    bot = CapBot()

    async def get_file(fid):
        return types.SimpleNamespace(file_path=fid)
    bot.get_file = get_file
    await T.run_op({"bot": bot, "redis": None}, job_id, CHAT, 9, "fa")
    async with maker() as s:
        job = await s.get(Job, job_id)
    assert job.status == "failed" and job.error == "pdf_needs_password"
    assert any(t("fa", "pdf_needs_password") in c for c in bot.captions), bot.captions


# ── تحویلِ آلبوم ─────────────────────────────────────────────────
class AlbumBot(ValidatingBot):
    def __init__(self, fail_groups: bool = False) -> None:
        self.groups: list[int] = []
        self.docs = 0
        self.fail_groups = fail_groups

    async def send_media_group(self, *a, **kw):
        payload = bind_like_aiogram("send_media_group", a, kw)
        if self.fail_groups:
            raise RuntimeError("group refused")
        self.groups.append(len(payload["media"]))
        return []

    def _on(self, name, payload):
        if name == "send_document":
            self.docs += 1
        return True


def _files(tmp_path, n):
    out = []
    for i in range(n):
        p = tmp_path / f"f{i}.pdf"
        p.write_bytes(b"%PDF-1.4\n")
        out.append(str(p))
    return out


async def test_album_is_sent_ten_at_a_time(tmp_path):
    bot = AlbumBot()
    failed, _ = await T._send_files(bot, CHAT, _files(tmp_path, 21), album=True)
    assert failed == 0 and bot.groups == [10, 10] and bot.docs == 1   # تنهای آخر: سندِ تک


async def test_a_refused_album_falls_back_to_single_files(tmp_path):
    bot = AlbumBot(fail_groups=True)
    failed, _ = await T._send_files(bot, CHAT, _files(tmp_path, 4), album=True)
    assert failed == 0 and bot.docs == 4


async def test_without_album_every_file_is_its_own_message(tmp_path):
    bot = AlbumBot()
    await T._send_files(bot, CHAT, _files(tmp_path, 3), album=False)
    assert bot.groups == [] and bot.docs == 3


# ── درِ ربات ─────────────────────────────────────────────────────
class RouterBot(ValidatingBot):
    def __init__(self) -> None:
        self.calls: list[str] = []

    def _on(self, name, payload):
        self.calls.append(name)
        return True


@pytest_asyncio.fixture
async def router_env(monkeypatch):
    engine = create_async_engine("sqlite+aiosqlite:///:memory:")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    maker = async_sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)
    notes: list[str] = []
    jobs: list[dict] = []

    async def note(bot, chat, mid, file, lang, note=None, **kw):
        bot.calls.append("note")
        notes.append(note or "")

    async def enqueue(bot, pool, session, file, op, args, *a, **kw):
        jobs.append({"op": op, **args})

    async def no_limit(pool, uid):
        return None
    monkeypatch.setattr(ops, "set_card_note", note)
    monkeypatch.setattr(ops, "_enqueue", enqueue)
    monkeypatch.setattr(ops, "_check_limits", no_limit)
    async with maker() as session:
        user = User(tg_user_id=777, role="user")
        session.add(user)
        await session.flush()
        f = File(ref="PdfRout1", owner_id=user.id, file_unique_id="u", file_id="f",
                 name="a.pdf", kind="pdf", size=10)
        session.add(f)
        await session.flush()
        state = FSMContext(storage=MemoryStorage(), key=StorageKey(bot_id=1, chat_id=CHAT, user_id=777))
        yield session, user, f, state, notes, jobs
    await engine.dispose()


def _text(bot, text: str, mid: int = 5) -> Message:
    return Message(message_id=mid, date=datetime.now(timezone.utc),
                   chat=Chat(id=CHAT, type="private"), text=text).as_(bot)


async def test_the_password_is_deleted_first_and_never_enters_the_job(router_env, redis):
    session, user, f, state, notes, jobs = router_env
    await state.set_state(PdfPassword.waiting)
    await state.update_data(ref=f.ref, card_chat=CHAT, card_mid=9, mode="lock")
    bot = RouterBot()
    await ops.op_pdf_pw_recv(_text(bot, "رمزِ من"), state, session, "fa", redis, user)
    assert bot.calls[0] == "delete_message", bot.calls
    assert jobs and jobs[0]["op"] == "pdf_lock" and set(jobs[0]) == {"op", "tok"}
    assert "رمزِ من" not in str(jobs)
    assert await redis.get(pagespec.PW_KEY.format(tok=jobs[0]["tok"])) == "رمزِ من"
    assert 0 < await redis.ttl(pagespec.PW_KEY.format(tok=jobs[0]["tok"])) <= pagespec.PW_TTL
    assert await state.get_state() is None


async def test_a_bad_password_stays_in_the_question(router_env, redis):
    session, user, f, state, notes, jobs = router_env
    await state.set_state(PdfPassword.waiting)
    await state.update_data(ref=f.ref, card_chat=CHAT, card_mid=9, mode="unlock")
    bot = RouterBot()
    await ops.op_pdf_pw_recv(_text(bot, "a\nb"), state, session, "fa", redis, user)
    assert bot.calls[0] == "delete_message"
    assert not jobs and notes[-1] == t("fa", "pdf_pw_bad")
    assert await state.get_state() == PdfPassword.waiting


async def test_page_numbers_are_checked_in_the_bot(router_env):
    session, user, f, state, notes, jobs = router_env
    await state.set_state(PdfPages.waiting)
    await state.update_data(ref=f.ref, card_chat=CHAT, card_mid=9, mode="delete")
    bot = RouterBot()
    await ops.op_pdf_pages_recv(_text(bot, "سه تا"), state, session, "fa", None, user)
    assert not jobs and notes[-1] == t("fa", "pdf_pages_bad")
    assert await state.get_state() == PdfPages.waiting
    await ops.op_pdf_pages_recv(_text(bot, "۲ تا ۴"), state, session, "fa", None, user)
    assert jobs == [{"op": "pdf_select", "spec": "۲ تا ۴", "mode": "delete"}]


def _doc_msg(bot, mid: int, uid: str) -> Message:
    return Message(message_id=mid, date=datetime.now(timezone.utc),
                   chat=Chat(id=CHAT, type="private"),
                   document=Document(file_id=f"d-{uid}", file_unique_id=uid, file_name=f"{uid}.pdf",
                                     mime_type="application/pdf", file_size=10)).as_(bot)


async def test_merge_order_is_the_send_order_not_the_arrival_order(router_env):
    """آلبومِ تلگرام هم‌زمان پردازش می‌شود؛ صفحهٔ ۳ می‌توانست پیش از ۲ ثبت شود."""
    session, user, f, state, notes, jobs = router_env
    await state.set_state(Collect.collecting)
    await state.update_data(ref=f.ref, card_chat=CHAT, card_mid=9, purpose="merge",
                            members=[{"file_id": "f", "name": "a.pdf", "size": 10, "mid": 0}])
    bot = RouterBot()
    await ops.collect_recv(_doc_msg(bot, 32, "third"), state, session, "fa", user)
    await ops.collect_recv(_doc_msg(bot, 31, "second"), state, session, "fa", user)
    names = [m["file_id"] for m in (await state.get_data())["members"]]
    assert names == ["f", "d-second", "d-third"]
