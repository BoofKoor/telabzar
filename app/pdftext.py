"""متنِ PDF با ترتیبِ درستِ راست‌به‌چپ — و دو خروجی از آن: TXT و Wordِ ویرایش‌پذیر.

**چرا نه فقط `pdftotext`.** poppler ترتیبِ خواندنِ فارسی را درست درمی‌آورد ولی
لیگاتورِ «لا» را **برعکس** باز می‌کند: «سلام» → «سالم»، «کلاس» → «کالس»، «بالا» →
«باال» (اندازه‌گیری‌شده روی PDFِ LibreOffice **و** Chrome، poppler 24.02). علتش این
است که poppler جعبهٔ یک گلیفِ چندحرفی را از چپ به راست بینِ حروفش تقسیم می‌کند و
بعد کلمهٔ راست‌به‌چپ را حرف‌به‌حرف برمی‌گرداند. در فارسی «لا» در کلمه‌های پرکاربرد
است، و چون «ال» هم کلمهٔ رایجی است، با پس‌پردازشِ رشته‌ای جبران‌پذیر نیست.

**و چرا نه فقط pdfplumber.** pdfminer هر گلیف را یک واحد نگه می‌دارد (لیگاتور سالم
می‌ماند) ولی ترتیبِ دیداری (چپ‌به‌راست) می‌دهد، و سطرِ مخلوطِ فارسی/انگلیسی را
براساسِ ترتیبِ جریانِ محتوا به چند سطرِ جدا می‌شکند؛ ستون‌بندی را هم نمی‌فهمد.

**پس ترکیبی:** ساختار (جریان → بلوک → سطر، و تشخیصِ ستون) از `pdftotext -bbox-layout`
و **محتوای** هر سطر از گلیف‌های pdfplumber که داخلِ جعبهٔ همان سطر می‌افتند. بعد
ترتیبِ دیداری با الگوریتمِ دوسطحیِ bidi (`vis2log`) به ترتیبِ منطقی برمی‌گردد —
گلیف به گلیف، پس «لا» یک واحد است و درونش برعکس نمی‌شود.

صفحه‌ای که هیچ حرفِ راست‌به‌چپ ندارد اصلاً pdfplumber نمی‌خواهد (کلمه‌های poppler
برای چپ‌به‌راست درست‌اند) — پس سندِ بزرگِ انگلیسی سریع می‌ماند. صفحهٔ بی‌متن
(اسکن) با OCR پر می‌شود، تا سقفِ صفحه‌ای که ادمین تعیین کرده.
"""
from __future__ import annotations

import asyncio
import logging
import os
import re
import statistics
import unicodedata
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field

from . import pdftools as T
from .exceptions import ProcessingCancelled, UserFacingError

log = logging.getLogger("telabzar.pdftext")

#: Word بیش از این صفحه نمی‌سازد (هر صفحه یک پارسِ کاملِ pdfminer است).
DOCX_MAX_PAGES = 200
#: TXT بیش از این صفحه نمی‌خواند — کتابِ هزارصفحه‌ای هم در چند دقیقه.
TXT_MAX_PAGES = 1000
#: عکسِ کوچک‌تر از این (pt) آیکون/خط/جداکننده است، نه تصویرِ محتوایی.
IMG_MIN_PT = 28
#: سقفِ تعدادِ عکس در یک خروجیِ Word.
IMG_MAX = 300
IMG_DPI = 150

_NS = "{http://www.w3.org/1999/xhtml}"


# ── bidi: ترتیبِ دیداری → منطقی ──────────────────────────────────
_R, _L, _NUM, _N = "R", "L", "#", None
_MIRROR = {"(": ")", ")": "(", "[": "]", "]": "[", "{": "}", "}": "{", "<": ">", ">": "<",
           "«": "»", "»": "«", "‹": "›", "›": "‹"}
#: پایانِ جمله — برای حدسِ جهتِ پاراگراف از جای نقطه (نگاه کن `block_rtl`)
_TERMINAL = set(".!?؟:;؛،,")


def _udir(unit: str) -> str | None:
    """جهتِ یک واحد (گلیف): اولین نویسهٔ قوی یا عددیِ آن."""
    for ch in unit:
        b = unicodedata.bidirectional(ch)
        if b in ("R", "AL"):
            return _R
        if b == "L":
            return _L
        if b in ("EN", "AN"):
            return _NUM
    return _N


def has_rtl(text: str) -> bool:
    return any(unicodedata.bidirectional(ch) in ("R", "AL") for ch in text or "")


def _levels(units: list[str], rtl: bool) -> list[int]:
    """سطحِ bidiِ هر واحد، روی ترتیبِ **دیداری** (UAX #9 در حدِ دو سطح).

    واحدهای همسایه در ترتیبِ دیداری همان همسایه‌های منطقی‌اند (فقط ترتیبِ گروه‌ها
    عوض می‌شود)، پس حلِ خنثی‌ها و عددها روی همین ترتیب همان نتیجه را می‌دهد.
    """
    n = len(units)
    dirs = [_udir(u) for u in units]

    def strong(i: int, step: int) -> str | None:
        j = i + step
        while 0 <= j < n:
            if dirs[j] in (_R, _L):
                return dirs[j]
            j += step
        return None

    # W4/W5 تقریبی: «۱۲:۳۰»، «۳.۵»، «٪۲۰» — جداکنندهٔ میانِ دو عدد خودش عدد است
    for i in range(1, n - 1):
        if dirs[i] is _N and dirs[i - 1] == _NUM and dirs[i + 1] == _NUM \
                and len(units[i]) == 1 and unicodedata.bidirectional(units[i]) in ("CS", "ES", "ET"):
            dirs[i] = _NUM
    # W7: عددی که قوی‌ترین همسایه‌اش چپ‌به‌راست است، خودش چپ‌به‌راست است
    for i in range(n):
        if dirs[i] == _NUM:
            a, b = strong(i, -1), strong(i, 1)
            if (a or b) and _R not in (a, b):
                dirs[i] = _L
    # N1/N2: خنثی میانِ دو هم‌جهت آن جهت را می‌گیرد (عدد مثلِ R)، وگرنه جهتِ پایه
    res: list[str] = []
    for i in range(n):
        d = dirs[i]
        if d is _N:
            a = strong_or_num(dirs, i, -1)
            b = strong_or_num(dirs, i, 1)
            d = a if (a is not None and a == b) else (_R if rtl else _L)
        res.append(d)
    out = []
    for d in res:
        if d == _R:
            out.append(1)
        elif d == _L:
            out.append(2 if rtl else 0)
        else:                       # عدد
            out.append(2)
    return out


def strong_or_num(dirs: list, i: int, step: int) -> str | None:
    """نزدیک‌ترین واحدِ غیرخنثی در یک جهت؛ عدد برای حلِ خنثی‌ها R حساب می‌شود (N1)."""
    j = i + step
    while 0 <= j < len(dirs):
        d = dirs[j]
        if d == _NUM:
            return _R
        if d in (_R, _L):
            return d
        j += step
    return None


def vis2log(units: list[str], rtl: bool) -> str:
    """ترتیبِ دیداریِ گلیف‌ها → متنِ منطقی.

    نمایش (UAX #9 L2) از بالاترین سطح تا پایین‌ترین سطحِ فرد، هر دنبالهٔ هم‌سطح یا
    بالاتر را برعکس می‌کند؛ هر برعکس‌کردن خودش وارونِ خودش است، پس معکوسِ کل همان
    برعکس‌کردن‌ها **به ترتیبِ برعکس** است: اول سطحِ ≥۱، بعد سطحِ ≥۲.
    گلیفِ سطحِ فرد در نمایش آینه می‌شود (L4)، پس این‌جا آینه برمی‌گردد — وگرنه
    «(با پرانتز)» به «)با پرانتز(» تبدیل می‌شد (اندازه‌گیری‌شده روی PDFِ Chrome).
    """
    if not units:
        return ""
    if not rtl and not any(_udir(u) == _R for u in units):
        return "".join(units)
    seq = list(zip(units, _levels(units, rtl)))
    top = max(lv for _, lv in seq)
    for k in range(1, top + 1):
        i = 0
        while i < len(seq):
            if seq[i][1] >= k:
                j = i
                while j < len(seq) and seq[j][1] >= k:
                    j += 1
                seq[i:j] = seq[i:j][::-1]
                i = j
            else:
                i += 1
    return "".join(_MIRROR.get(u, u) if lv % 2 else u for u, lv in seq)


def block_rtl(lines_units: list[list[str]]) -> bool:
    """جهتِ پایهٔ یک پاراگراف از همهٔ سطرهایش — نه سطربه‌سطر، وگرنه سطرِ تمام‌انگلیسیِ
    وسطِ پاراگرافِ فارسی جهتِ خودش را می‌گرفت و نقطه‌ها جابه‌جا می‌شدند.

    اکثریتِ حروفِ قوی تصمیم می‌گیرد، با یک استثنا که از دادهٔ واقعی آمد: پاراگرافِ
    **چپ‌به‌راستی** که متنش فارسی است (خروجیِ LibreOffice از یک TXT) نقطهٔ پایانی را
    سمتِ **راست** دارد؛ در پاراگرافِ راست‌به‌چپ همان نقطه سمتِ **چپ** می‌نشیند.
    پس جای علامتِ پایانِ جمله در دو سرِ سطرها رأیِ اکثریت را می‌شکند.
    """
    r = l = 0
    left_end = right_end = 0
    for units in lines_units:
        for u in units:
            d = _udir(u)
            if d == _R:
                r += 1
            elif d == _L:
                l += 1
        core = [u for u in units if u.strip()]
        if len(core) >= 2:
            if core[0] in _TERMINAL and core[-1] not in _TERMINAL:
                left_end += 1
            elif core[-1] in _TERMINAL and core[0] not in _TERMINAL:
                right_end += 1
    if r == 0:
        return False
    if l == 0:
        return not (right_end and not left_end)
    if left_end != right_end:
        return left_end > right_end
    return r >= l


def logical_rtl(text: str) -> bool:
    """جهتِ پاراگرافِ متنِ **منطقی** (متنِ فایلِ TXT، خروجیِ tesseract): اکثریتِ حروفِ قوی.

    با `block_rtl` اشتباه گرفته نشود: آن ترتیبِ **دیداریِ** گلیف‌های PDF را می‌گیرد و
    «نقطه سمتِ راست» را نشانهٔ پاراگرافِ چپ‌به‌راست می‌خواند. در متنِ منطقی نقطهٔ پایانِ
    جملهٔ فارسی هم در **انتهای** رشته است، پس همان قاعده هر جملهٔ فارسیِ نقطه‌دار را
    چپ‌چین می‌کرد (اندازه‌گیری‌شده روی TXT → PDF).
    """
    r = l = 0
    for ch in text or "":
        b = unicodedata.bidirectional(ch)
        if b in ("R", "AL"):
            r += 1
        elif b == "L":
            l += 1
    return r > l


# ── ساختارِ صفحه از poppler ──────────────────────────────────────
@dataclass
class Line:
    x0: float
    y0: float
    x1: float
    y1: float
    flow: int = 0                                     # ستون/جریانِ poppler
    words: list[str] = field(default_factory=list)   # متنِ poppler (ترتیبِ دیداری)
    word_w: list[float] = field(default_factory=list)  # پهنای همان کلمه‌ها (pt)
    units: list[str] = field(default_factory=list)   # گلیف‌های pdfplumber (دیداری)
    sizes: list[float] = field(default_factory=list)
    bold: int = 0
    text: str = ""

    @property
    def size(self) -> float:
        s = [x for x in self.sizes if x > 0]
        return statistics.median(s) if s else (self.y1 - self.y0) / 1.2


@dataclass
class Para:
    """یک پاراگراف: سطرهای پشتِ‌همِ یک جریان، هم‌اندازه و بی‌فاصلهٔ اضافه."""
    lines: list[Line]
    rtl: bool = False
    center: bool = False

    @property
    def x0(self) -> float:
        return min(ln.x0 for ln in self.lines)

    @property
    def x1(self) -> float:
        return max(ln.x1 for ln in self.lines)

    @property
    def y0(self) -> float:
        return self.lines[0].y0

    @property
    def size(self) -> float:
        s = [x for ln in self.lines for x in ln.sizes if x > 0]
        return statistics.median(s) if s else statistics.median([ln.size for ln in self.lines])

    @property
    def bold(self) -> bool:
        glyphs = sum(len(ln.sizes) for ln in self.lines)
        return glyphs > 0 and sum(ln.bold for ln in self.lines) > glyphs / 2

    def joined(self) -> str:
        """سطرها پشتِ‌هم برای Word (که خودش می‌شکندشان). «exam-» + «ple» → «example»."""
        buf: list[str] = []
        for ln in self.lines:
            t = ln.text
            if not t:
                continue
            if buf and not self.rtl and buf[-1].endswith("-") and t[:1].islower():
                buf[-1] = buf[-1][:-1] + t
            else:
                buf.append(t)
        return " ".join(buf)


@dataclass
class Img:
    x0: float
    y0: float
    x1: float
    y1: float
    path: str = ""


@dataclass
class Page:
    no: int
    width: float
    height: float
    lines: list[Line] = field(default_factory=list)   # به ترتیبِ خواندنِ poppler
    paras: list[Para] = field(default_factory=list)
    images: list[Img] = field(default_factory=list)
    ocr: str | None = None          # متنِ OCRِ صفحهٔ اسکن‌شده
    scan: bool = False              # بی‌متن و (برای Word) بی‌OCR → عکسِ کلِ صفحه


@dataclass
class Doc:
    pages: list[Page]
    total: int                      # کلِ صفحه‌های PDF
    ocr_pages: int = 0
    ocr_skipped: int = 0            # صفحهٔ بی‌متنی که از سقفِ OCR جا ماند


def _f(el, key: str) -> float:
    return float(el.get(key, "0"))


def parse_bbox(path: str) -> list[Page]:
    """خروجیِ `pdftotext -bbox-layout` → صفحه/سطر (با شمارهٔ جریان). iterparse تا سندِ
    بزرگ کلِ درختِ XML را در حافظه نسازد.

    «بلوک»های poppler عمداً دور ریخته می‌شوند: روی صفحهٔ دوستونیِ Chrome هر سطر یک
    بلوکِ جدا بود (اندازه‌گیری‌شده)، پس بلوک پاراگراف نیست. پاراگراف از هندسهٔ
    سطرها ساخته می‌شود (`group`)؛ چیزی که از poppler می‌ماند ترتیبِ خواندن و
    **جریان** است — همان که ستون‌ها را از هم جدا می‌کند.
    """
    pages: list[Page] = []
    cur: Page | None = None
    ln: Line | None = None
    flow = -1
    for ev, el in ET.iterparse(path, events=("start", "end")):
        tag = el.tag.replace(_NS, "")
        if ev == "start":
            if tag == "page":
                cur = Page(no=len(pages) + 1, width=_f(el, "width"), height=_f(el, "height"))
            elif tag == "flow":
                flow += 1
            elif tag == "line" and cur is not None:
                ln = Line(_f(el, "xMin"), _f(el, "yMin"), _f(el, "xMax"), _f(el, "yMax"), flow=flow)
            continue
        if tag == "word" and ln is not None:
            ln.words.append(el.text or "")
            ln.word_w.append(_f(el, "xMax") - _f(el, "xMin"))
        elif tag == "line" and cur is not None and ln is not None:
            if ln.words:
                cur.lines.append(ln)
            ln = None
        elif tag == "page" and cur is not None:
            pages.append(cur)
            cur = None
            el.clear()
    return pages


# ── محتوای سطر از گلیف‌های pdfplumber ────────────────────────────
def _is_mark(text: str) -> bool:
    return bool(text) and all(unicodedata.category(ch) in ("Mn", "Me") for ch in text)


def _joining(text: str) -> bool:
    return any(unicodedata.bidirectional(ch) == "AL" for ch in text)


def line_units(chars: list[dict]) -> tuple[list[str], list[float], int]:
    """گلیف‌های یک سطر (دیکشنریِ pdfplumber) → (واحدهای دیداری, اندازه‌ها, شمارِ پررنگ).

    سه ظرافت، هر سه از PDFهای واقعی:
    * **فاصله** یا گلیفِ فاصلهٔ صریح است یا شکافی بزرگ‌تر از ربعِ اندازهٔ قلم.
    * **نیم‌فاصله**: LibreOffice نویسهٔ ZWNJ را با یک گلیفِ «فاصله» می‌نویسد که پیشروی
      ندارد و رویِ حرفِ بعدی می‌افتد — «پی‌دی‌اف» در pdfplumber «پی دی اف» می‌شد.
      فاصلهٔ صریحی که میانِ دو حرفِ پیوسته‌نویس شکافی ایجاد نکرده، ZWNJ است.
    * **اِعراب** (فتحه، تنوین، …) جعبه‌ای درونِ حرفِ پایه دارد؛ اگر واحدِ جدا بماند
      بعد از برگرداندنِ ترتیب پیش از حرفش می‌نشیند. به حرفِ زیرینش چسبانده می‌شود.
    """
    glyphs = [c for c in chars if c["text"] and not c["text"].isspace() and not _is_mark(c["text"])]
    spaces = [c for c in chars if c["text"] and c["text"].isspace()]
    marks = [c for c in chars if _is_mark(c["text"])]
    glyphs.sort(key=lambda c: (c["x0"], c["x1"]))
    texts = [unicodedata.normalize("NFKC", c["text"]) for c in glyphs]
    for m in marks:                                   # اِعراب → حرفِ زیرین
        cx = (m["x0"] + m["x1"]) / 2
        best, dist = None, None
        for i, g in enumerate(glyphs):
            d = 0.0 if g["x0"] <= cx <= g["x1"] else min(abs(cx - g["x0"]), abs(cx - g["x1"]))
            if dist is None or d < dist:
                best, dist = i, d
        if best is not None:
            texts[best] += m["text"]
    units: list[str] = []
    sizes: list[float] = []
    bold = 0
    for i, g in enumerate(glyphs):
        if i:
            prev = glyphs[i - 1]
            gap = g["x0"] - prev["x1"]
            size = max(g.get("size") or 0, prev.get("size") or 0, 1.0)
            # فاصله‌ای که **از** شکافِ همین دو گلیف شروع می‌شود. نسخهٔ گشادترِ این شرط
            # (مرکزِ فاصله تا نیم‌قلم بعد از گلیف) فاصلهٔ کلمهٔ **بعدی** را هم می‌گرفت و
            # در «پرانتز» نیم‌فاصلهٔ ساختگی و در «(با» فاصلهٔ اضافه می‌گذاشت.
            between = [s for s in spaces if prev["x1"] - 0.5 <= s["x0"] <= g["x0"] + 0.5]
            if between and gap < size * 0.12 and _joining(texts[i - 1]) and _joining(texts[i]):
                units.append("‌")
            elif between or gap > size * 0.25:
                units.append(" ")
        units.append(texts[i])
        sizes.append(float(g.get("size") or 0))
        if "bold" in (g.get("fontname") or "").lower():
            bold += 1
    return units, sizes, bold


def _assign(page: Page, chars: list[dict]) -> None:
    """هر گلیف → سطری از poppler که مرکزش در جعبهٔ آن است (نزدیک‌ترین، اگر چند تا)."""
    lines = page.lines
    buckets: list[list[dict]] = [[] for _ in lines]
    # نمایهٔ عمودی: هر سطر در خانه‌های یک‌نقطه‌ایِ ارتفاعِ خودش. بدونِ آن هر گلیف همهٔ
    # سطرهای صفحه را می‌گشت — روی صفحهٔ پر یک‌سومِ کلِ زمانِ استخراج.
    rows: dict[int, list[int]] = {}
    for i, ln in enumerate(lines):
        for y in range(int(ln.y0 - 1), int(ln.y1 + 1) + 1):
            rows.setdefault(y, []).append(i)
    for c in chars:
        cx, cy = (c["x0"] + c["x1"]) / 2, (c["top"] + c["bottom"]) / 2
        best, dist = -1, None
        for i in rows.get(int(cy), ()):
            ln = lines[i]
            if ln.y0 - 1 <= cy <= ln.y1 + 1 and ln.x0 - 2 <= cx <= ln.x1 + 2:
                d = abs(cy - (ln.y0 + ln.y1) / 2)
                if dist is None or d < dist:
                    best, dist = i, d
        if best >= 0:
            buckets[best].append(c)
    for ln, cs in zip(lines, buckets):
        if any(not c["text"].isspace() for c in cs):
            ln.units, ln.sizes, ln.bold = line_units(cs)


def _line_rtl(ln: Line) -> bool:
    return block_rtl([ln.units])


def _overlap(a: Line, b: Line) -> float:
    return min(a.x1, b.x1) - max(a.x0, b.x0)


def _extents(lines: list[Line]) -> list[tuple[float, float]]:
    """لبه‌های ستونِ هر سطر: کم‌ترین x0 و بیش‌ترین x1 در میانِ سطرهایی که با آن
    هم‌ستون‌اند (بیش از نیمی از سطرِ باریک‌تر رویش افتاده). «جریان»ِ poppler این‌جا
    به کار نمی‌آید: روی PDFِ Chrome سطرِ اولِ پاراگراف و بقیه‌اش دو جریانِ جدا بودند،
    پس سطرِ کوتاهِ آخرِ پاراگراف «پهنای ستونِ» خودش را یک سطر می‌دید."""
    out = []
    for a in lines:
        same = [b for b in lines
                if _overlap(a, b) > 0.5 * min(a.x1 - a.x0, b.x1 - b.x0)]
        out.append((min(b.x0 for b in same), max(b.x1 for b in same)))
    return out


def group(page: Page) -> list[Para]:
    """سطرها (به ترتیبِ خواندن) → پاراگراف. پاراگرافِ تازه وقتی:
    * ستون عوض شود: هم‌پوشانیِ افقی با سطرِ قبل کمتر از نیمِ سطرِ باریک‌تر، یا سطر
      **بالاتر** از قبلی بیفتد (ترتیبِ خواندن به بالای ستونِ بعد پریده)؛
    * فاصلهٔ عمودی از ۶۰٪ ارتفاعِ سطر بیشتر شود (فاصلهٔ میانِ پاراگراف‌ها)؛
    * اندازهٔ قلم بیش از ۱۵٪ عوض شود (تیتر → متن)؛
    * سطرِ قبلی «کوتاه» تمام شده باشد (پایینِ همین تابع)؛
    * جهتِ سطر عوض شود.
    """
    paras: list[Para] = []
    ext = _extents(page.lines)
    cur: list[Line] = []
    prev_i = -1
    for i, ln in enumerate(page.lines):
        if cur:
            prev = cur[-1]
            fx0, fx1 = ext[prev_i]
            fw = max(fx1 - fx0, 1.0)
            h = max((prev.y1 - prev.y0 + ln.y1 - ln.y0) / 2, 1.0)
            gap = ln.y0 - prev.y1
            ps, ls = prev.size, ln.size
            prev_rtl = _line_rtl(prev)
            # «کوتاه» یعنی کلمهٔ اولِ سطرِ بعد در جای خالیِ این سطر **جا می‌شد** و
            # نویسنده باز هم سطر را شکسته — پس پایانِ پاراگراف است. آستانهٔ ثابت (۱۵٪
            # پهنا) متنِ بی‌ترازِ چپ‌چین را سطر به سطر می‌شکست.
            room = (prev.x0 - fx0) if prev_rtl else (fx1 - prev.x1)
            first_w = (ln.word_w[-1] if prev_rtl else ln.word_w[0]) if ln.word_w else ps * 2
            short = room > max(first_w + ps * 0.5, fw * 0.05)
            moved = _overlap(prev, ln) < 0.5 * min(prev.x1 - prev.x0, ln.x1 - ln.x0) \
                or ln.y1 < prev.y0
            if (moved or gap > h * 0.6 or abs(ps - ls) > max(ps, ls) * 0.15 or short
                    or _line_rtl(ln) != prev_rtl):
                paras.append(Para(cur))
                cur = []
        cur.append(ln)
        prev_i = i
    if cur:
        paras.append(Para(cur))
    return paras


def _columns(band: list[Line]) -> list[list[Line]]:
    """سطرهای یک نوار → ستون‌ها (دسته‌های هم‌پوشانِ افقی)."""
    cols: list[list[Line]] = []
    for ln in band:
        hit = [c for c in cols if any(min(ln.x1, o.x1) - max(ln.x0, o.x0) > 0 for o in c)]
        merged = [ln]
        for c in hit:
            merged.extend(c)
            cols.remove(c)
        cols.append(merged)
    return cols


def _rtl_column_order(lines: list[Line]) -> list[Line]:
    """ستون‌های راست‌به‌چپ از ستونِ **راست** خوانده می‌شوند.

    poppler ستون‌ها را همیشه از چپ می‌خواند، و «جریان»هایش هم ستون نیستند: روی صفحهٔ
    دوستونیِ فارسی (Chrome با `dir=rtl`) تیترِ بالای راست و کلِ ستونِ **چپ** یک جریان
    شدند و پاراگرافِ سوم پیش از اول درآمد (اندازه‌گیری‌شده). پس:

    * سطرِ «پهن» (بیش از ۵۵٪ پهنای متن — تیترِ سرتاسری، پاراگرافِ تک‌ستونه) و هر
      نوارِ سفیدِ افقیِ بلندتر از نیم‌سطر صفحه را به نوارهای افقی می‌بُرند — وگرنه
      بخشِ دوستونیِ انگلیسیِ بالای صفحه و بخشِ دوستونیِ فارسیِ پایین یک نوار می‌شدند
      و رأیِ اکثریت را انگلیسی می‌برد؛
    * سطرهای باریکِ هر نوار با هم‌پوشانیِ افقی به ستون دسته می‌شوند؛
    * فقط نواری که **دست‌کم دو ستونِ سه‌سطری** دارد و بیشترِ سطرهایش راست‌به‌چپ‌اند
      از راست به چپ بازچیده می‌شود، هر ستون با `flow`ِ تازه تا `group` پاراگراف را سرِ
      ستون ببندد.

    هر چیزِ دیگر — صفحهٔ چپ‌به‌راست، صفحهٔ تک‌ستونهٔ فارسی، نوارِ تک‌ستونه — **همان
    ترتیب و جریانِ poppler** را نگه می‌دارد؛ برای آن‌ها ترتیبِ poppler درست است و
    بازچینی فقط پاراگرافی را که از یک نوار به نوارِ بعد ادامه دارد می‌شکست.
    """
    if len(lines) < 6:
        return lines
    index = {id(ln): i for i, ln in enumerate(lines)}
    left, right = min(ln.x0 for ln in lines), max(ln.x1 for ln in lines)
    wide = (right - left) * 0.55
    gap = statistics.median([ln.y1 - ln.y0 for ln in lines]) * 0.5
    bands: list[list[Line]] = []
    cur: list[Line] = []
    bottom = None
    for ln in sorted(lines, key=lambda x: (x.y0, -x.x1)):
        if ln.x1 - ln.x0 > wide:
            if cur:
                bands.append(cur)
                cur = []
            bands.append([ln])
            bottom = None
            continue
        if cur and bottom is not None and ln.y0 - bottom > gap:
            bands.append(cur)          # نوارِ سفیدِ افقی: بخشِ تازه (مثلاً تیترِ بخش)
            cur = []
        cur.append(ln)
        bottom = ln.y1 if bottom is None or not cur[:-1] else max(bottom, ln.y1)
    if cur:
        bands.append(cur)
    plan: list[tuple[list[Line], list[list[Line]] | None]] = []
    changed = False
    for band in bands:
        cols = _columns(band)
        rtl = sum(1 for ln in band if any(has_rtl(w) for w in ln.words)) > len(band) / 2
        if rtl and sum(1 for c in cols if len(c) >= 3) >= 2:
            plan.append((band, sorted(cols, key=lambda c: -max(x.x1 for x in c))))
            changed = True
        else:
            plan.append((band, None))
    if not changed:
        return lines
    out: list[Line] = []
    fresh = max(ln.flow for ln in lines) + 1
    for band, cols in plan:
        if cols is None:
            out.extend(sorted(band, key=lambda x: index[id(x)]))
            continue
        for c in cols:
            for ln in sorted(c, key=lambda x: x.y0):
                ln.flow = fresh
                out.append(ln)
            fresh += 1
    return out


def _finish_page(page: Page) -> None:
    """پاراگراف‌ها، جهتِ هرکدام، و متنِ منطقیِ هر سطر. سطری که گلیفِ pdfplumber نگرفته
    (صفحهٔ چرخیده، یا صفحهٔ بی‌حرفِ راست‌به‌چپ) به کلمه‌های poppler برمی‌گردد."""
    for ln in page.lines:
        if not ln.units:
            joined = " ".join(unicodedata.normalize("NFKC", w) for w in ln.words)
            ln.units = list(joined)
    page.lines = _rtl_column_order(page.lines)
    page.paras = group(page)
    for para in page.paras:
        para.rtl = block_rtl([ln.units for ln in para.lines])
        for ln in para.lines:
            ln.text = re.sub(r" {2,}", " ", vis2log(ln.units, para.rtl)).strip()
        if page.width:
            mid = (para.x0 + para.x1) / 2
            para.center = (abs(mid - page.width / 2) < page.width * 0.04
                           and para.x1 - para.x0 < page.width * 0.7 and len(para.lines) <= 3)


def _page_needs_glyphs(page: Page) -> bool:
    return any(has_rtl(w) for ln in page.lines for w in ln.words)


def dedupe(chars: list[dict]) -> list[dict]:
    """گلیفِ تکراری (پررنگِ ساختگی: همان حرف دو بار با جابه‌جاییِ کسرِ نقطه) یک‌بار.

    `Page.dedupe_chars`ِ خودِ pdfplumber روی ۶۰ صفحهٔ فارسیِ پر **۱۷ ثانیه از ۳۰**
    را می‌خورد (مرتب‌سازی‌های تودرتو، اندازه‌گیری‌شده)؛ این نسخهٔ خطی همان کار را با
    یک مجموعه می‌کند. کلید: متن + اندازه + مختصاتِ گردشده به یک نقطه.
    """
    seen: set[tuple] = set()
    out: list[dict] = []
    for c in chars:
        key = (c["text"], round(c["x0"]), round(c["top"]), round(c.get("size") or 0))
        if key in seen:
            continue
        seen.add(key)
        out.append(c)
    return out


def _plumb(src: str, pages: list[Page], want_images: bool, cancel_flag) -> None:
    """در thread: گلیف‌ها (و برای Word جعبهٔ عکس‌ها) از pdfplumber، صفحه به صفحه."""
    import pdfplumber  # ورودِ تنبل: فقط ورکر این وابستگی را دارد

    with pdfplumber.open(src) as pdf:
        for page in pages:
            if cancel_flag():
                return
            need = _page_needs_glyphs(page)
            if not need and not want_images:
                continue
            try:
                pp = pdf.pages[page.no - 1]
                # صفحهٔ چرخیده: مختصاتِ دو ابزار ممکن است هم‌خوان نباشد → فقط poppler
                same = abs(pp.width - page.width) < 2 and abs(pp.height - page.height) < 2
                if need and same:
                    _assign(page, dedupe(pp.chars))
                if want_images and same:
                    page.images = _page_images(pp.images, page)
                pp.close()
            except Exception:  # noqa: BLE001 — صفحهٔ خراب نباید کلِ سند را بشکند
                log.warning("pdfplumber failed on page %s", page.no, exc_info=True)


def _page_images(imgs: list[dict], page: Page) -> list[Img]:
    """عکس‌های محتوایی: نه آیکونِ ریز، نه تکراری، و نه پس‌زمینهٔ تمام‌صفحه روی صفحه‌ای
    که متن دارد (اسکنِ OCRشده یا تصویرِ زمینه — متنش از قبل استخراج شده)."""
    out: list[Img] = []
    seen: set[tuple] = set()
    area = page.width * page.height or 1
    for im in imgs:
        x0, y0 = max(0.0, im["x0"]), max(0.0, im["top"])
        x1, y1 = min(page.width, im["x1"]), min(page.height, im["bottom"])
        w, h = x1 - x0, y1 - y0
        if w < IMG_MIN_PT or h < IMG_MIN_PT:
            continue
        if page.lines and w * h > area * 0.8:
            continue
        key = (round(x0), round(y0), round(x1), round(y1))
        if key in seen:
            continue
        seen.add(key)
        out.append(Img(x0, y0, x1, y1))
    return out


async def extract(src: str, workdir: str, *, max_pages: int, ocr_max_pages: int,
                  want_images: bool, cancel=None, progress=None) -> Doc:
    """`src` باید خروجیِ `pdftools.prepare` باشد (بی‌قفل، PDFِ واقعی)."""
    total = await T.page_count(src, cancel=cancel)
    last = min(total, max_pages)
    bbox = os.path.join(workdir, "_bbox.html")
    rc, _, err = await T._capture(
        ["pdftotext", "-bbox-layout", "-enc", "UTF-8", "-f", "1", "-l", str(last), src, bbox],
        timeout=600, cancel=cancel)
    if rc != 0 or not os.path.exists(bbox):
        raise UserFacingError("pdf_damaged", detail=T._tail(err))
    pages = await asyncio.to_thread(parse_bbox, bbox)
    # poppler صفحهٔ کاملاً خالی را هم یک <page> می‌نویسد؛ اگر نشمرد، جای خالی پر شود
    while len(pages) < last:
        pages.append(Page(no=len(pages) + 1, width=0, height=0))

    # گلیف‌ها در thread؛ لغو بینِ صفحه‌ها پرسیده می‌شود
    cancelled = {"v": False}

    async def _watch() -> None:
        while cancel is not None:
            await asyncio.sleep(2)
            try:
                if await cancel():
                    cancelled["v"] = True
                    return
            except Exception:  # noqa: BLE001
                pass

    watcher = asyncio.create_task(_watch())
    try:
        await asyncio.to_thread(_plumb, src, pages, want_images, lambda: cancelled["v"])
    finally:
        watcher.cancel()
    if cancelled["v"]:
        raise ProcessingCancelled()
    for p in pages:
        _finish_page(p)

    doc = Doc(pages=pages, total=total)
    empty = [p for p in pages if not p.lines]
    ocr_these = empty[:max(0, ocr_max_pages)] if T.have(T.TESSERACT) else []
    doc.ocr_skipped = len(empty) - len(ocr_these)
    if ocr_these:
        texts = await T.ocr_pages(src, workdir, [p.no for p in ocr_these],
                                  cancel=cancel, progress=progress)
        for p, txt in zip(ocr_these, texts):
            p.ocr = T.tidy_ocr(txt)
        doc.ocr_pages = len(ocr_these)
    for p in empty:
        if p.ocr is None:
            p.scan = True
    return doc


# ── خروجیِ TXT ───────────────────────────────────────────────────
_LRM, _RLM = "‎", "‏"


def _mark_dir(text: str, rtl: bool) -> str:
    """نمایشگرِ متنِ ساده جهتِ پاراگراف را از **اولین** حرفِ قوی حدس می‌زند؛ سطرِ
    فارسی‌ای که با کلمهٔ انگلیسی شروع شود را چپ‌به‌راست می‌چید و نقطه‌اش جابه‌جا
    می‌شد. یک نشانهٔ نامرئیِ جهت فقط همان سطرها را درست می‌کند."""
    first = next((d for d in map(_udir, text) if d in (_R, _L)), None)
    if rtl and first == _L:
        return _RLM + text
    if not rtl and first == _R:
        return _LRM + text
    return text


def to_text(doc: Doc) -> str:
    """سطرها همان سطرهای PDF؛ میانِ پاراگراف‌ها یک سطرِ خالی، میانِ صفحه‌ها دو سطر."""
    pages: list[str] = []
    for p in doc.pages:
        if p.ocr is not None:
            if p.ocr:
                pages.append("\n".join(_mark_dir(ln, has_rtl(ln)) if ln else ""
                                       for ln in p.ocr.split("\n")))
            continue
        paras = ["\n".join(_mark_dir(ln.text, para.rtl) for ln in para.lines if ln.text)
                 for para in p.paras]
        body = "\n\n".join(x for x in paras if x)
        if body:
            pages.append(body)
    return "\n\n\n".join(pages).strip() + "\n"


def has_text(doc: Doc) -> bool:
    return any(p.lines or p.ocr for p in doc.pages)


# ── خروجیِ Word ──────────────────────────────────────────────────
def _runs(text: str, rtl: bool) -> list[tuple[str, bool]]:
    """متنِ یک پاراگراف → تکه‌های هم‌جهت. Word تکهٔ `w:rtl` را با قلم/اندازهٔ
    «پیچیده» (cs) می‌چیند؛ انگلیسیِ وسطِ پاراگرافِ فارسی تکهٔ جدای خودش است."""
    out: list[list] = []
    for ch in text:
        d = _udir(ch)
        is_r = (d == _R) if d in (_R, _L) else None
        if not out:
            out.append([ch, rtl if is_r is None else is_r])
        elif is_r is None or is_r == out[-1][1]:
            out[-1][0] += ch
        else:
            out.append([ch, is_r])
    return [(t, r) for t, r in out]


def _set_rtl_para(p) -> None:
    from docx.oxml import OxmlElement

    p_pr = p._p.get_or_add_pPr()
    p_pr.insert(0, OxmlElement("w:bidi"))


def _style_run(run, size: float, bold: bool, rtl: bool) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    run.font.size = Pt(size)
    if bold:
        run.bold = True
    r_pr = run._r.get_or_add_rPr()
    sz_cs = OxmlElement("w:szCs")
    sz_cs.set(qn("w:val"), str(int(round(size * 2))))
    r_pr.append(sz_cs)
    if bold:
        r_pr.append(OxmlElement("w:bCs"))
    if rtl:
        r_pr.append(OxmlElement("w:rtl"))


async def render_images(src: str, workdir: str, doc: Doc, cancel=None) -> None:
    """هر عکس = همان ناحیه از صفحهٔ رندرشده — نه بایت‌های خامِ جریانِ تصویر، که ماسک،
    فضای رنگ و چرخشِ PDF را نمی‌شناسند. صفحه یک‌بار رندر و عکس‌ها از آن برش می‌خورند.
    صفحهٔ اسکنِ بی‌OCR (`scan`) خودش یک عکسِ تمام‌صفحه می‌شود."""
    from PIL import Image

    count = 0
    for p in doc.pages:
        targets = list(p.images)
        if p.scan:
            targets = [Img(0, 0, p.width, p.height)]
        if not targets or count >= IMG_MAX or not p.width:
            p.images = []
            continue
        png = await T.render_page(src, p.no, os.path.join(workdir, f"_pg-{p.no}"),
                                  dpi=IMG_DPI, cancel=cancel)
        kept: list[Img] = []
        with Image.open(png) as im:
            sx, sy = im.width / p.width, im.height / p.height
            for k, img in enumerate(targets):
                if count >= IMG_MAX:
                    break
                box = (int(img.x0 * sx), int(img.y0 * sy), int(img.x1 * sx), int(img.y1 * sy))
                if box[2] - box[0] < 4 or box[3] - box[1] < 4:
                    continue
                out = os.path.join(workdir, f"_img-{p.no}-{k}.jpg")
                im.crop(box).convert("RGB").save(out, "JPEG", quality=85)
                img.path = out
                kept.append(img)
                count += 1
        os.remove(png)
        p.images = kept


def read_text_file(path: str) -> str:
    """فایلِ متنیِ کاربر → رشته. UTF-8 (با/بی BOM) و UTF-16ِ BOM‌دار، وگرنه
    Windows-1256 — فایلِ متنیِ فارسیِ قدیمی (Notepadِ ویندوزِ فارسی) همین است و
    LibreOffice آن را به حروفِ لاتینِ درهم باز می‌کرد."""
    with open(path, "rb") as fh:
        raw = fh.read()
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16", "replace")
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        return raw.decode("cp1256", "replace")


def text_to_docx(text: str, out: str, font: str = "Noto Sans Arabic") -> None:
    """متنِ ساده → Word با جهتِ درستِ **هر پاراگراف** (برای TXT → PDF).

    LibreOffice فایلِ TXT را با پاراگرافِ چپ‌به‌راست باز می‌کند، پس متنِ فارسی چپ‌چین
    و با نقطهٔ سمتِ راست چاپ می‌شد (اندازه‌گیری‌شده). این‌جا هر سطر یک پاراگراف است
    و جهتش از حروفِ خودش. قلمِ «پیچیده» Noto Sans Arabic است (`fonts-noto-core`ِ ایمیجِ
    ورکر)؛ اگر نبود LibreOffice از fontconfig قلمِ فارسی‌دارِ دیگری برمی‌دارد.
    """
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt

    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)   # A4، نه Letterِ پیش‌فرضِ python-docx
    normal = d.styles["Normal"]
    normal.font.size = Pt(11)
    fonts = normal.element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:cs"), font)
    fonts.set(qn("w:ascii"), "DejaVu Sans")
    fonts.set(qn("w:hAnsi"), "DejaVu Sans")
    normal.paragraph_format.space_after = Pt(2)
    for line in (text or "").replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        line = line.rstrip()
        par = d.add_paragraph()
        if not line:
            continue
        rtl = logical_rtl(line)
        if rtl:
            _set_rtl_para(par)
        for chunk, r in _runs(line, rtl):
            _style_run(par.add_run(chunk), 11, False, r)
    d.save(out)


def to_docx(doc: Doc, out: str) -> None:
    """سندِ Wordِ ویرایش‌پذیر: پاراگرافِ جاری (نه کادرهای شناورِ LibreOffice)، جهتِ
    درست، اندازه و پررنگیِ تقریبی، عکس‌ها سرِ جایشان، شکستِ صفحه میانِ صفحه‌ها."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    d = Document()
    paras = [para for p in doc.pages for para in p.paras]
    sizes = [x for para in paras for ln in para.lines for x in ln.sizes if x > 0]
    body = statistics.median(sizes) if sizes else 11.0
    if paras:
        rtl_doc = sum(1 for x in paras if x.rtl) > len(paras) / 2
    else:
        rtl_doc = any(has_rtl(p.ocr or "") for p in doc.pages)

    # قلمِ «پیچیده» (فارسی) Tahoma: روی هر ویندوزی هست و فارسیِ خوانا دارد؛ قلمِ
    # اصلیِ PDF (B Nazanin و …) ممکن است روی دستگاهِ کاربر نباشد.
    normal = d.styles["Normal"]
    normal.font.size = Pt(round(body * 2) / 2)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:cs"), "Tahoma")

    first = next((p for p in doc.pages if p.width), None)
    sec = d.sections[0]
    if first is not None:
        sec.page_width, sec.page_height = Pt(first.width), Pt(first.height)
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, m, Pt(48))
    if rtl_doc:
        sec._sectPr.append(OxmlElement("w:bidi"))
    text_w = (first.width if first else 595) - 96

    def add_text(text: str, rtl: bool, size: float, bold: bool, center: bool) -> None:
        par = d.add_paragraph()
        if rtl:
            _set_rtl_para(par)
        if center:
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for chunk, r in _runs(text, rtl):
            _style_run(par.add_run(chunk), round(size * 2) / 2, bold, r)

    def add_image(img: Img) -> None:
        if not img.path:
            return
        try:
            d.add_picture(img.path, width=Pt(min(img.x1 - img.x0, text_w)))
        except Exception:  # noqa: BLE001 — یک عکسِ خراب نباید کلِ سند را بشکند
            log.warning("docx picture failed: %s", img.path, exc_info=True)
            return
        pic = d.paragraphs[-1]
        pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if rtl_doc:
            _set_rtl_para(pic)

    for i, p in enumerate(doc.pages):
        if i:
            d.add_page_break()
        if p.ocr is not None:
            for chunk in re.split(r"\n\s*\n", p.ocr):
                text = " ".join(x.strip() for x in chunk.splitlines() if x.strip())
                if text:
                    add_text(text, logical_rtl(text), body, False, False)
            continue
        pending = sorted(p.images, key=lambda im: im.y0)
        for para in p.paras:
            # عکسی که بالای این پاراگراف است **و** با آن هم‌ستون (هم‌پوشانیِ افقی) پیش از
            # آن می‌آید. ترتیبِ پاراگراف‌ها ترتیبِ خواندنِ poppler است (ستون‌به‌ستون)،
            # پس عکسِ ستونِ کناری جلوی پاراگرافی از ستونِ دیگر نمی‌پرد.
            ready = [im for im in pending
                     if im.y0 <= para.y0 and min(im.x1, para.x1) - max(im.x0, para.x0) > 0]
            for im in ready:
                pending.remove(im)
                add_image(im)
            text = para.joined()
            if text:
                add_text(text, para.rtl, para.size, para.bold, para.center)
        for img in pending:
            add_image(img)
    d.save(out)
